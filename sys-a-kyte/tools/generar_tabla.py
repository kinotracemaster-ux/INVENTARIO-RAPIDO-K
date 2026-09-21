#!/usr/bin/env python3
"""Genera app/tabla.js (la tabla base de la app) desde un Excel con hojas en formato Kyte.

Uso:
    python3 tools/generar_tabla.py [excel] [salida]

Por defecto lee datos/INVENTARIO_KINO_18_de_sep.xlsx y escribe app/tabla.js.

Qué hace (igual que el botón "Cambiar tabla" de la app):
  - Toma cada hoja que tenga las columnas de Kyte (ID, Nombre*, Código, Precio*, Stock actual).
  - Se queda con las 14 columnas del formato Kyte, en orden. Ignora la columna auxiliar "SYS".
  - Usa los valores guardados (no fórmulas) y normaliza saltos de línea a \\n.
  - Recorta espacios en el nombre de la hoja (ej. "POEDAGAR " -> "POEDAGAR").
"""
import sys, json, datetime, pathlib, unicodedata, re
import openpyxl

ROOT = pathlib.Path(__file__).resolve().parent.parent
COLS = ['ID', 'Nombre*', 'Categoría', 'Código', 'Descripción', 'Costo unitario', 'Precio*',
        'Precio de promoción', 'Mostrar en el catálogo', 'Destacar', 'Fracción?',
        'Controlar stock', 'Stock actual', 'Stock mínimo']
REQ = ['ID', 'Nombre*', 'Código', 'Precio*', 'Stock actual']
INT_COLS = (6, 12, 13)  # Precio*, Stock actual, Stock mínimo: enteros si vienen como 5.0

src = pathlib.Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / 'datos' / 'INVENTARIO_KINO_18_de_sep.xlsx'
out = pathlib.Path(sys.argv[2]) if len(sys.argv) > 2 else ROOT / 'app' / 'tabla.js'


def norm(s):
    s = unicodedata.normalize('NFD', str(s if s is not None else ''))
    s = ''.join(c for c in s if not unicodedata.combining(c)).lower()
    return re.sub(r'[^a-z0-9]+', ' ', s).strip()


wb = openpyxl.load_workbook(src, data_only=True)
hojas = []
for ws in wb.worksheets:
    hdr = [norm(ws.cell(1, c).value) for c in range(1, ws.max_column + 1)]
    idx = [hdr.index(norm(c)) if norm(c) in hdr else -1 for c in COLS]
    if any(idx[COLS.index(r)] < 0 for r in REQ):
        continue  # no es una hoja de Kyte (ej. SYS INV)
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        row = [None if i < 0 else r[i] for i in idx]
        if all(v is None or v == '' for v in row):
            continue
        row = [v.replace('\r\n', '\n').replace('\r', '\n') if isinstance(v, str) else v for v in row]
        row = [int(v) if isinstance(v, float) and v.is_integer() and j in INT_COLS else v
               for j, v in enumerate(row)]
        rows.append(row)
    hojas.append({'nombre': ws.title.strip(), 'rows': rows})
    print(f'  {ws.title.strip():20s} {len(rows):6d} productos')

if not hojas:
    sys.exit('No encontré hojas con formato Kyte en ese Excel.')

tabla = {'v': 1, 'guardada': datetime.datetime.now().astimezone().isoformat(timespec='seconds'),
         'origen': src.name, 'cols': COLS, 'hojas': hojas}
out.write_text('window.__TABLA__=' + json.dumps(tabla, ensure_ascii=False, separators=(',', ':')) + ';',
               encoding='utf-8')
print('escrito', out, f'({out.stat().st_size:,} bytes)')
