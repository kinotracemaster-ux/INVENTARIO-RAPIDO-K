import os, re, json, base64, threading, functools, io, zipfile, sys, pathlib
from http.server import ThreadingHTTPServer, SimpleHTTPRequestHandler
import openpyxl
from playwright.sync_api import sync_playwright

SP = str(pathlib.Path(__file__).resolve().parent.parent)  # raíz del proyecto
SRC = os.environ.get('SYS_XLSX') or f'{SP}/datos/INVENTARIO_KINO_18_de_sep.xlsx'
OUT = f'{SP}/test/out'
os.makedirs(OUT, exist_ok=True)
COLS = ['ID','Nombre*','Categoría','Código','Descripción','Costo unitario','Precio*','Precio de promoción','Mostrar en el catálogo','Destacar','Fracción?','Controlar stock','Stock actual','Stock mínimo']
FAILS = []
def check(cond, msg):
    print(('PASS ' if cond else 'FAIL ') + msg)
    if not cond: FAILS.append(msg)

# ---------- Resultado esperado, calculado aparte (Python) ----------
def kk1(s): return str(s).strip().upper()
def kk2(s): return re.sub(r'[\s\-_:./\\]+', '', str(s).upper())
def blank(v): return v is None or v == ''
def load_expected(mult=2, price=True, sys_path=SRC, table_path=SRC):
    wbs = openpyxl.load_workbook(sys_path, data_only=True)
    ws = wbs['SYS INV']
    k1, k2, lista = {}, {}, []
    for art, desc, stock, precio in ws.iter_rows(min_row=2, max_col=4, values_only=True):
        if art is None: continue
        a = str(art).strip()
        if a.upper() in k1: continue
        rec = dict(art=a, stock=stock, precio=precio)
        k1[a.upper()] = rec; lista.append(rec)
        k2.setdefault(kk2(a), []).append(rec)
    wbt = openpyxl.load_workbook(table_path, data_only=True)
    res = {}
    used = set()
    for name in ['KYTE CON SYS INV', 'CURREN', 'JOEFOX', 'POEDAGAR ']:
        w = wbt[name]
        hdr = [w.cell(1, c).value for c in range(1, w.max_column + 1)]
        idx = [hdr.index(c) for c in COLS]
        rows_in, rows_out = [], []
        st = dict(prod=0, cruzan=0, sd=0, pd=0, sin=0, formato=0)
        for r in w.iter_rows(min_row=2, values_only=True):
            row = [r[i] for i in idx]
            if all(blank(v) for v in row): continue
            st['prod'] += 1
            out = list(row)
            cod = '' if blank(row[3]) else str(row[3]).strip()
            rec = None; tier = ''
            if cod:
                rec = k1.get(cod.upper()); tier = 'exacto'
                if rec is None:
                    a = k2.get(kk2(cod))
                    if a and len(a) == 1: rec = a[0]; tier = 'formato'
            if rec is None:
                st['sin'] += 1
            else:
                used.add(id(rec)); st['cruzan'] += 1
                if tier == 'formato': st['formato'] += 1
                antes = None if blank(row[12]) else row[12]
                despues = max(0, rec['stock'])
                if antes is None and despues == 0: despues = None
                if despues != antes: st['sd'] += 1
                out[12] = despues
                if price and rec['precio'] and rec['precio'] > 0:
                    nuevo = round(rec['precio'] * mult)
                    if nuevo != row[6]: st['pd'] += 1; out[6] = nuevo
            rows_in.append(row); rows_out.append(out)
        res[name.strip()] = dict(rows=rows_out, st=st)
    nuevos = [x for x in lista if id(x) not in used]
    res['_nuevos_total'] = len(nuevos)
    res['_nuevos_con_stock'] = len([x for x in nuevos if x['stock'] > 0])
    return res

def read_xlsx_rows(data, sheet=None):
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    return wb.sheetnames, [[c for c in r] for r in ws.iter_rows(values_only=True)]

def same(a, b):
    if blank(a) and blank(b): return True
    if isinstance(a, (int, float)) and isinstance(b, (int, float)): return abs(a - b) < 1e-9
    return a == b

def compare_rows(got, exp):
    """got: filas leídas (con encabezado). exp: filas esperadas (sin encabezado)"""
    if [str(x) for x in got[0]] != COLS: return False, f'encabezado distinto: {got[0]}'
    got = got[1:]
    if len(got) != len(exp): return False, f'filas {len(got)} != {len(exp)}'
    for i, (g, e) in enumerate(zip(got, exp)):
        for j in range(14):
            if not same(g[j], e[j]):
                return False, f'fila {i+2} col {COLS[j]}: obtuve {g[j]!r}, esperaba {e[j]!r}'
    return True, ''

# ---------- Servidor local + navegador ----------
class Q(SimpleHTTPRequestHandler):
    def log_message(self, *a): pass
srv = ThreadingHTTPServer(('127.0.0.1', 0), functools.partial(Q, directory=f'{SP}/app'))
port = srv.server_address[1]
threading.Thread(target=srv.serve_forever, daemon=True).start()

STUB = """
window.__saved=[]; window.__published=[];
function b64(u8){let s='';const CH=0x8000;for(let i=0;i<u8.length;i+=CH)s+=String.fromCharCode.apply(null,u8.subarray(i,i+CH));return btoa(s);}
window.claude={use: async (n)=>{
  if(n==='downloads') return {save: async (req)=>{
    let d=req.data, u8;
    if(d instanceof Uint8Array) u8=d; else if(d instanceof ArrayBuffer) u8=new Uint8Array(d); else if(d instanceof Blob) u8=new Uint8Array(await d.arrayBuffer()); else u8=new TextEncoder().encode(String(d));
    window.__saved.push({f:req.filename,b:b64(u8)}); return {status:'saved'};
  }};
  if(n==='artifact') return {publish: async (files)=>{ window.__published.push(files); return {version:'v2'}; }};
  return null;
}};
"""

def run_flow(p):
    browser = p.chromium.launch(args=['--no-sandbox'])
    ctx = browser.new_context(viewport={'width': 1200, 'height': 900}, accept_downloads=True)
    ctx.add_init_script(STUB)
    ctx.route('https://cdn.jsdelivr.net/npm/xlsx@0.18.5/dist/xlsx.full.min.js', lambda r: r.fulfill(path=f'{SP}/test/node_modules/xlsx/dist/xlsx.full.min.js', content_type='application/javascript'))
    ctx.route('https://cdn.jsdelivr.net/npm/jszip@3.10.1/dist/jszip.min.js', lambda r: r.fulfill(path=f'{SP}/test/node_modules/jszip/dist/jszip.min.js', content_type='application/javascript'))
    ctx.route(re.compile(r'https://fonts\.(googleapis|gstatic)\.com/.*'), lambda r: r.abort())
    page = ctx.new_page()
    errors = []
    page.on('pageerror', lambda e: errors.append(str(e)))
    page.on('console', lambda m: errors.append('console.' + m.type + ': ' + m.text) if m.type in ('error',) and 'fonts' not in m.text and 'Failed to load resource' not in m.text else None)
    page.goto(f'http://127.0.0.1:{port}/index.html')
    page.wait_for_selector('#chain .node')

    # 1) Estado inicial con la tabla real
    nodes = page.locator('#chain article.node')
    check(nodes.count() == 4, 'hay 4 nodos de hojas (KYTE CON, CURREN, JOEFOX, POEDAGAR)')
    inicial = [n.locator('.big').inner_text() for n in nodes.all()]
    print('   conteos iniciales:', inicial)
    check(inicial == ['5.792', '294', '245', '478'], 'conteos iniciales de la tabla base')
    check('18 sep 2026' in page.locator('#base').inner_text(), 'muestra la fecha de la tabla base')
    check(page.locator('#btn-zip').is_disabled(), 'zip deshabilitado antes de subir el SYS')

    def descargar(selector, n_esperado):
        page.click(selector)
        page.wait_for_function(f'window.__saved.length>={n_esperado}', timeout=30000)
        page.wait_for_function('!document.querySelector("#btn-zip").disabled', timeout=30000)
    def guardados():
        out = {}
        for s in page.evaluate('()=>window.__saved'):
            data = base64.b64decode(s['b']); out[s['f']] = data
            open(f"{OUT}/{s['f']}", 'wb').write(data)
        return out

    # 2) Subir el SYS real (el libro completo; debe usar la hoja 'SYS INV')
    page.set_input_files('#f-sys', SRC)
    page.wait_for_selector('.node.sys.on', timeout=30000)
    exp = load_expected()
    for k in ['KYTE CON SYS INV', 'CURREN', 'JOEFOX', 'POEDAGAR']:
        print('   esperado', k, exp[k]['st'])
    sys_txt = page.locator('#dz').inner_text()
    print('   nodo SYS:', sys_txt.replace('\n', ' | '))
    check('9.993' in sys_txt, 'lee 9.993 artículos del SYS')
    check('3.140' in sys_txt, 'cuenta 3.140 con stock')
    ui_sd = [n.locator('.big').inner_text() for n in nodes.all()]
    exp_sd = [f"{exp[k]['st']['sd']:,}".replace(',', '.') for k in ['KYTE CON SYS INV', 'CURREN', 'JOEFOX', 'POEDAGAR']]
    print('   UI cambian de stock:', ui_sd, '| esperado:', exp_sd)
    check(ui_sd == exp_sd, 'conteo "cambian de stock" por hoja coincide con el cálculo aparte')
    facts = page.locator('#facts').inner_text().replace('\n', ' | ')
    print('   facts:', facts)
    kst = exp['KYTE CON SYS INV']['st']
    check(f"{kst['cruzan']:,}".replace(',', '.') in facts and f"{kst['prod']:,}".replace(',', '.') in facts, 'facts: cruzan / productos')
    check(f"{exp['_nuevos_total']:,}".replace(',', '.') in facts, 'facts: nuevos en SYS')

    # 3) Descargar 4 hojas + zip + reporte
    names = ['KYTE CON SYS INV', 'CURREN', 'JOEFOX', 'POEDAGAR']
    for i in range(4):
        descargar(f'#btn-bajar-{i}', i + 1)
    descargar('#btn-zip', 5)
    descargar('#btn-reporte', 6)
    files = guardados()
    print('   archivos:', {k: len(v) for k, v in files.items()})
    check(len(files) == 6, 'se entregaron 6 archivos (4 hojas, zip, reporte)')
    for i, nm in enumerate(names):
        fn = [f for f in files if f.startswith(nm.replace(' ', '_') + '_')]
        check(len(fn) == 1, f'archivo de {nm} presente')
        sheets, rows = read_xlsx_rows(files[fn[0]])
        ok, why = compare_rows(rows, exp[nm]['rows'])
        check(ok, f'{nm}: TODAS las celdas coinciden con lo esperado ({len(rows)-1} filas) {why}')
        check(sheets == [nm], f'{nm}: nombre de hoja {sheets}')
    z = [f for f in files if f.endswith('.zip')][0]
    zf = zipfile.ZipFile(io.BytesIO(files[z]))
    check(len(zf.namelist()) == 4, f'zip trae 4 archivos: {zf.namelist()}')
    for nm in names:
        zn = [n for n in zf.namelist() if n.startswith(nm.replace(' ', '_') + '_')][0]
        _, rows = read_xlsx_rows(zf.read(zn))
        ok, why = compare_rows(rows, exp[nm]['rows'])
        check(ok, f'zip/{nm} coincide con lo esperado {why}')
    rep = [f for f in files if f.startswith('Reporte_SYS')][0]
    wbr = openpyxl.load_workbook(io.BytesIO(files[rep]), data_only=True)
    print('   hojas del reporte:', {ws.title: ws.max_row for ws in wbr.worksheets})
    check(wbr.sheetnames == ['Resumen', 'Cambios', 'Sin cruzar', 'Nuevos en SYS', 'Por formato'], f'hojas del reporte: {wbr.sheetnames}')
    check(wbr['Sin cruzar'].max_row - 1 == kst['sin'], f"reporte: sin cruzar = {kst['sin']}")
    check(wbr['Nuevos en SYS'].max_row - 1 == exp['_nuevos_con_stock'], f"reporte: nuevos con stock = {exp['_nuevos_con_stock']}")
    check(wbr['Por formato'].max_row - 1 == kst['formato'], f"reporte: por formato = {kst['formato']}")
    check(wbr['Cambios'].max_row - 1 >= max(kst['sd'], kst['pd']), 'reporte: hoja Cambios con filas')

    # 4) Vista previa de cambios
    page.wait_for_selector('#lista table')
    n_rows = page.locator('#lista tbody tr').count()
    check(n_rows == 100, f'vista previa muestra 100 filas primero ({n_rows})')
    page.click('#btn-mas')
    check(page.locator('#lista tbody tr').count() == 300, 'mostrar más agrega 200')
    page.fill('#buscar', '1401-3')
    txt = page.locator('#lista').inner_text()
    check('1401-3' in txt, 'búsqueda por código funciona')
    page.fill('#buscar', '')
    page.click('[data-act="chip"][data-f="precio"]')
    filas_precio = page.locator('#lista tbody tr').count()
    check(filas_precio == min(100, kst['pd']), f'filtro Precio: {filas_precio} filas visibles (esperado {min(100,kst["pd"])})')
    page.click('[data-act="chip"][data-f="todos"]')

    # captura de escritorio (una sola vez)
    page.screenshot(path=f'{OUT}/desktop.png', full_page=False)

    # 5) Apagar precio => solo stock
    page.evaluate('()=>{window.__saved.length=0}')
    page.uncheck('#opt-precio')
    page.wait_for_function('document.querySelectorAll("#chain article.node .sm")[0].innerText.startsWith("0 de precio")', timeout=15000)
    descargar('#btn-bajar-0', 1)
    files2 = guardados()
    expn = load_expected(price=False)
    fn = [f for f in files2 if f.startswith('KYTE_CON_SYS_INV_')][0]
    _, rows = read_xlsx_rows(files2[fn])
    ok, why = compare_rows(rows, expn['KYTE CON SYS INV']['rows'])
    check(ok, f'sin actualizar precio: KYTE CON coincide (precios intactos) {why}')
    page.check('#opt-precio')
    page.wait_for_function('!document.querySelectorAll("#chain article.node .sm")[0].innerText.startsWith("0 de precio")', timeout=15000)

    # 6) Cambiar tabla con el Excel original: lo leído debe ser igual a tabla.js
    page.set_input_files('#f-tabla', SRC)
    page.wait_for_selector('#btn-pend-si', timeout=30000)
    print('   pendiente:', page.locator('#pend').inner_text().replace('\n', ' | '))
    page.click('#btn-pend-si')
    page.wait_for_function('window.__published.length>=1', timeout=15000)
    pub = page.evaluate('()=>window.__published[0]')
    check(list(pub.keys()) == ['tabla.js'], 'guardar tabla publica solo tabla.js')
    js = pub['tabla.js']
    nueva = json.loads(js[len('window.__TABLA__='):-1])
    orig_js = open(f'{SP}/app/tabla.js', encoding='utf-8').read()
    orig = json.loads(orig_js[len('window.__TABLA__='):-1])
    check([h['nombre'] for h in nueva['hojas']] == [h['nombre'] for h in orig['hojas']], f"hojas leídas: {[h['nombre'] for h in nueva['hojas']]}")
    todo_igual = True; motivo = ''
    for hn, ho in zip(nueva['hojas'], orig['hojas']):
        if len(hn['rows']) != len(ho['rows']): todo_igual = False; motivo = f"{hn['nombre']}: {len(hn['rows'])} vs {len(ho['rows'])}"; break
        for i, (a, b) in enumerate(zip(hn['rows'], ho['rows'])):
            for j in range(14):
                if not same(a[j], b[j]): todo_igual = False; motivo = f"{hn['nombre']} fila {i} col {j}: {a[j]!r} vs {b[j]!r}"; break
            if not todo_igual: break
        if not todo_igual: break
    check(todo_igual, f'la tabla leída del Excel es idéntica a tabla.js {motivo}')
    page.wait_for_selector('#toast:not([hidden])')
    check('Tabla guardada' in page.locator('#toast').inner_text(), 'aviso: tabla guardada')

    # 7) Errores: archivo que no es SYS
    wbx = openpyxl.Workbook(); wbx.active.append(['foo', 'bar']); wbx.active.append([1, 2]); wbx.save(f'{OUT}/no_es_sys.xlsx')
    page.set_input_files('#f-sys', f'{OUT}/no_es_sys.xlsx')
    page.wait_for_selector('.aviso.err', timeout=10000)
    check('No encontré las columnas' in page.locator('#aviso').inner_text(), 'archivo equivocado: mensaje claro')

    # 8) SYS distinto: uno con muy pocos cruces dispara la alerta
    wb2 = openpyxl.Workbook(); ws2 = wb2.active
    ws2.append(['Artículo', 'Descripción', 'Exist. unidades', 'Precio'])
    for i in range(50): ws2.append([f'ZZ-{i}', 'X', 5, 1000])
    ws2.append(['1401-3', 'Poedagar', 7, 130000])
    wb2.save(f'{OUT}/sys_pocos.xlsx')
    page.set_input_files('#f-sys', f'{OUT}/sys_pocos.xlsx')
    page.wait_for_selector('.aviso.warn', timeout=15000)
    check('Solo cruzó' in page.locator('#aviso').inner_text(), 'SYS con pocos cruces: alerta de archivo equivocado')

    # 9) CSV (coma y punto y coma)
    import csv
    wbs = openpyxl.load_workbook(SRC, data_only=True)['SYS INV']
    for sep, nm in [(',', 'sys_coma.csv'), (';', 'sys_puntocoma.csv')]:
        with open(f'{OUT}/{nm}', 'w', newline='', encoding='utf-8-sig') as fh:
            w = csv.writer(fh, delimiter=sep)
            for r in wbs.iter_rows(values_only=True): w.writerow(r)
        page.set_input_files('#f-sys', f'{OUT}/{nm}')
        page.wait_for_function('document.querySelector("#dz") && document.querySelector("#dz").innerText.includes("'+nm+'")', timeout=30000)
        nS = page.locator('#dz').inner_text().replace('\n', ' | ')
        print('   CSV', repr(sep), '->', nS)
        check('9.993' in nS, f'CSV con separador {sep!r}: lee 9.993 artículos')

    print('ERRORES DE PÁGINA:', errors if errors else 'ninguno')
    check(not errors, 'sin errores de JavaScript en la consola')

    # 10) Móvil (una captura)
    page.set_input_files('#f-sys', SRC)
    page.wait_for_selector('.node.sys.on', timeout=30000)
    page.wait_for_function('document.querySelector("#dz").innerText.includes("INVENTARIO")', timeout=30000)
    page.set_viewport_size({'width': 400, 'height': 860})
    page.screenshot(path=f'{OUT}/mobile.png', full_page=True)
    ov = page.evaluate('()=>document.documentElement.scrollWidth>document.documentElement.clientWidth')
    check(not ov, 'móvil: sin scroll horizontal de la página')
    browser.close()

with sync_playwright() as p:
    run_flow(p)
print('\nRESUMEN:', 'TODO OK' if not FAILS else f'{len(FAILS)} FALLAS -> {FAILS}')
